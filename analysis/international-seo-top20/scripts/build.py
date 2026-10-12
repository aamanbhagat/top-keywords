import sys, os, re, json
import pandas as pd
B=os.path.dirname(os.path.abspath(__file__)); S=os.path.dirname(B)
sys.path.insert(0,B)
import candidates, details, serp_evidence, discovery
OUT='/home/user/top-keywords/analysis/international-seo-top20'
os.makedirs(OUT,exist_ok=True)
master=pd.read_csv(os.path.join(S,'master.csv'))
audit=pd.read_csv(os.path.join(S,'screenshot_audit.csv'))
VR={'0 – 10':0,'10 – 100':1,'100 – 1K':2,'1K – 10K':3,'10K – 100K':4,'100K – 1M':5,'1M – 10M':6}
LB={0:0,1:10,2:100,3:1000,4:10000,5:100000,6:1000000}; UB={0:10,1:100,2:1000,3:10000,4:100000,5:1000000,6:10000000}
serp={s['serp_id']:s for s in serp_evidence.S}
# ---- scores & ranking
rows=[]
for c in candidates.C:
    c['total']=sum(c[k] for k in 'ABCDEFG')
cands=sorted(candidates.C,key=lambda x:(-x['total'],-x['A'],-x['F'],-x['B'],-x['E']))
rank=0
for c in cands:
    if not c.get('reject'): rank+=1; c['rank']=rank
    else: c['rank']=None
assert rank==20, rank
# ---- clusters
def members(c):
    parts=[]
    for mk,inc,exc in c['spec']:
        x=master[(master.market==mk)&master.keyword.str.contains(inc,case=False,regex=True)]
        if exc: x=x[~x.keyword.str.contains(exc,case=False,regex=True)]
        parts.append(x)
    return pd.concat(parts).drop_duplicates(['market','keyword'])
memrows=[]
for c in cands:
    x=members(c)
    v=x[x.vol_rank.notna()]
    c['kw_in_csv']=len(x); c['kw_visible']=int(x.in_screenshot.sum()); c['kw_with_volume']=len(v)
    c['n_10k_plus']=int((v.vol_rank>=4).sum()); c['n_1k_10k']=int((v.vol_rank==3).sum()); c['n_100_1k']=int((v.vol_rank==2).sum())
    c['floor']=int(sum(LB[int(k)] for k in v.vol_rank)); c['ceiling']=int(sum(UB[int(k)] for k in v.vol_rank))
    c['not_visible']=int((~x.in_screenshot).sum())
    p=master[(master.keyword==c['primary'])&(master.market.isin(c['markets'].split(',')))]
    if len(p):
        p=p.iloc[0]
        def f(v): return '' if pd.isna(v) or v=='' else v
        c['p_vol']=f(p.vol) or ('— (unknown)' if p.in_screenshot else 'not visible'); c['p_m3']=f(p.m3); c['p_yoy']=f(p.yoy); c['p_comp']=f(p.comp)
        c['p_bid']=('' if pd.isna(p.bid_lo) else f'₹{p.bid_lo:,.2f}–₹{p.bid_hi:,.2f}') ; c['p_src']=f(p.screenshot_files)
    else:
        c['p_vol']=c['p_m3']=c['p_yoy']=c['p_comp']=c['p_bid']=c['p_src']=''
    top=v.sort_values(['vol_rank','keyword'],ascending=[False,True]).head(8)
    c['top_kw']='; '.join(f'{t.keyword} ({t.vol}{", "+t.comp if isinstance(t.comp,str) and t.comp!="Low" else ""})' for t in top.itertuples())
    bids=v.bid_hi.dropna(); c['max_bid_hi_inr']=float(bids.max()) if len(bids) else None
    c['comp_mix']=', '.join(f'{k}:{n}' for k,n in v.comp.fillna('—').value_counts().items())
    if c['rank']:
        for t in x.sort_values(['vol_rank','keyword'],ascending=[False,True],na_position='last').itertuples():
            memrows.append(dict(rank=c['rank'],opportunity_id=c['id'],market=t.market,country=t.country,language=t.language,keyword=t.keyword,
                avg_monthly_searches=t.vol if isinstance(t.vol,str) and t.vol else ('— (shown as dash = unknown)' if t.in_screenshot else 'NOT VISIBLE IN SCREENSHOTS'),
                three_month_change=t.m3 if isinstance(t.m3,str) else '', yoy_change=t.yoy if isinstance(t.yoy,str) else '',
                ads_competition=t.comp if isinstance(t.comp,str) else '', top_of_page_bid_low_INR=t.bid_lo, top_of_page_bid_high_INR=t.bid_hi,
                measurement_period='Sep 2025 – Aug 2026', kp_location=('Japan' if t.market in('JP','JP_EN') else t.country), kp_language='All languages',
                source_csv=t.source_file, source_screenshot=t.screenshot_files if isinstance(t.screenshot_files,str) else '', validation_status=t.validation_status))
mem=pd.DataFrame(memrows)
CONF={'BR-01':'Medium-High'}
def serp_urls(c):
    out=[]
    for sid in c['serp'].split(';'):
        s=serp.get(sid)
        if s: out.append(f"{sid} [{s['search_mode']}] \"{s['query']}\": "+' | '.join(s['urls'][:6]))
    return '\n'.join(out)
def serp_assess(c):
    return ' / '.join(f"{sid}: {serp[sid]['assessment']}" for sid in c['serp'].split(';') if sid in serp)
# ---- Top20 table
t20=[]
for c in [c for c in cands if c['rank']]:
    d=details.D[c['id']]
    t20.append(dict(rank=c['rank'],id=c['id'],short_name=details.SHORT[c['id']],concept=c['concept'],markets=c['markets'],language=c['lang'],primary_keyword=c['primary'],
        primary_avg_monthly_searches=c['p_vol'],primary_3m_change=c['p_m3'],primary_yoy=c['p_yoy'],primary_ads_competition=c['p_comp'],primary_top_bid_INR=c['p_bid'],primary_source_screenshot=c['p_src'],
        cluster_kw_in_csv=c['kw_in_csv'],cluster_kw_with_volume=c['kw_with_volume'],cluster_10k_plus=c['n_10k_plus'],cluster_1k_10k=c['n_1k_10k'],cluster_100_1k=c['n_100_1k'],
        cluster_floor_sum=c['floor'],cluster_ceiling_sum=c['ceiling'],cluster_not_visible=c['not_visible'],max_top_bid_high_INR=c['max_bid_hi_inr'],
        related_keywords=c['top_kw'],
        A_organic_25=c['A'],B_demand_20=c['B'],C_toolfit_20=c['C'],D_intent_10=c['D'],E_cluster_10=c['E'],F_monetisation_10=c['F'],G_feasibility_5=c['G'],total_100=c['total'],
        confidence=c['confidence'],organic_assessment=c['A_why'],serp_evidence=serp_assess(c),serp_urls=serp_urls(c),competitors_and_weaknesses=d['comp'],
        why_this_one=d['why'],url_structure=d['url'],mvp=d['mvp'],localisation=d['loc'],monetisation=d['money'],
        score_explanation=f"A {c['A']}: {c['A_why']} | B {c['B']}: {c['B_why']} | C {c['C']}: {c['C_why']} | D {c['D']}: {c['D_why']} | E {c['E']}: {c['E_why']} | F {c['F']}: {c['F_why']} | G {c['G']}: {c['G_why']}",
        validate_before_build=d['validate']))
T20=pd.DataFrame(t20)
allc=pd.DataFrame([dict(rank=c['rank'] or '',id=c['id'],markets=c['markets'],concept=c['concept'],primary_keyword=c['primary'],primary_vol=c['p_vol'],
    kw_with_volume=c['kw_with_volume'],n_10k_plus=c['n_10k_plus'],n_1k_10k=c['n_1k_10k'],n_100_1k=c['n_100_1k'],cluster_floor_sum=c['floor'],
    A=c['A'],B=c['B'],C=c['C'],D=c['D'],E=c['E'],F=c['F'],G=c['G'],total=c['total'],organic_gate_pass=c['A']>=8,
    decision=('TOP 20' if c['rank'] else 'REJECTED'),reason=c.get('reject',''),confidence=c['confidence'],serp_ids=c['serp']) for c in cands])
# ---- SERP sheet
serp_df=pd.DataFrame([dict(serp_id=s['serp_id'],market=s['market'],query=s['query'],search_mode=s['search_mode'],date=s['date'],
    candidates=s['candidates'],assessment=s['assessment'],urls='\n'.join(s['urls']),n_urls=len(s['urls']),
    method_note='WebSearch tool (US index) with native-language query — proxy, not a geo-located Google SERP') for s in serp_evidence.S])
disc=pd.DataFrame(discovery.L,columns=['market','keyword_or_family_native','english_gloss','source_how_found','why_it_matters','validation_status'])
# ---- rubric
rubric=pd.DataFrame([
 ('A','Organic ranking opportunity & SERP weakness',25,'19–23: ≤1 dedicated tool, no brand/official tool | 14–18: several small tools, no brand, documented accuracy/coverage gaps | 10–13: many (≥6) small tools or exact-match domains, or one mid-authority brand | 8–9: calculator networks / strong specialist present | <8: brand- or official-tool-dominated → GATE: excluded from Top 20. Capped at 23 because all SERP evidence is proxy (not geo-located). Google Ads competition is NOT used here.'),
 ('B','Search demand & evidence reliability',20,'Demand 0–14 from Keyword Planner buckets of the cluster (14: ≥3 keywords ≥10K–100K; 11–12: one 10K–100K head + tail; 9–10: ≥4 at 1K–10K; 7–8: 2–3 at 1K–10K; 4–6: one 1K–10K; 1–3: only ≤1K; 0: no data) + reliability 0–6 (6 visible & visually verified; −1 head year-tagged/decaying or −90% YoY; 2–3 KP location mismatch; 0 no KP data).'),
 ('C','Fit for a tool/calculator/converter/lookup',20,'18–20 deterministic and fully answers the query in-page; 15–17 lookup/linguistic engine with edge cases; 11–14 estimator under discretion or partly informational; ≤10 informational/news or content-heavy.'),
 ('D','Search intent & relevance to site',10,'9–10 utility intent; 6–8 mixed; ≤5 informational/commercial-product/navigational.'),
 ('E','Traffic potential across cluster',10,'Breadth of validated keywords, programmatic/sibling pages, evergreen or recurring demand.'),
 ('F','Monetisation potential',10,'9–10 B2B/legal/finance lead-gen or high top-of-page bids; 6–8 consumer finance/property/insurance; 3–5 ads + light affiliates; 1–2 ads only in low-RPM markets.'),
 ('G','Implementation feasibility & localisation',5,'5 static rules, low liability; 3–4 annual updates or linguistic work; 1–2 frequent regulatory churn, legal liability or heavy data.'),
 ('Tie-break','',0,'Total → A → F → B → E.'),
 ('Penalty','',0,'Missing organic evidence never counts as low difficulty: unverifiable SERPs capped at A≤10; no KP data → reliability 0.'),
],columns=['code','criterion','max_points','rules'])
# ---- roadmap & validation
roadmap=pd.DataFrame([
 ('Batch 0 — Foundation','Weeks 0–2','Platform','Locale-folder architecture with hreflang (/pt-br/, /es-mx/, /es-es/, /id/, /ja/, /en/, /pl/, /de/, /tr/); shared calculator UI kit; parameter registry (value, effective date, source URL); analytics events; Search Console per folder; schema (WebApplication/FAQ/HowTo); Core Web Vitals budget.','All later batches','Medium'),
 ('Batch 1 — Fast, low-liability, high-evidence','Weeks 2–6','MX-01, BR-01, MX-03, ID-01, ID-09, JP-TXT','MX-01 first (aguinaldo season peaks Nov–Dec 2026 — target live by mid-Nov); BR-01 + land-unit engine; Banxico-backed MX-03; Javanese-calendar engine shared by ID-01 and ID-09 (publish kalender jawa 2027 before December); JP text toolkit (client-side).','Batch 0; MX-01 needs labour-law review','Low–Medium per tool (2–6 dev-days each)'),
 ('Batch 2 — Linguistic/data engines','Weeks 6–12','BR-03, PL-02, NUM, JP-01, PL-05','CTN finder (official data ingestion), Polish morphology (shared by PL-02 & NUM), JP address romanisation with Japan Post data, word finder (needs licensed word list).','Data licences (SGJP/word list, Japan Post), BR-03 correlation source','Medium–High'),
 ('Batch 3 — Regulated B2B calculators','Weeks 10–20','DE-02, DE-07, DE-03, PL-04, JP-07, BR-02, BR-09','Need professional review and update discipline; DE legal trio shares one "Recht & Kosten" hub; BR-02 builds on BR-03 search engine; BR-09 needs fee-schedule monitoring.','Expert review; parameter registry; geo-SERP validation (highest uncertainty group)','Medium–High'),
 ('Batch 4 — Timed / niche','Dec 2026 – Jan 2027','BR-08, PL-08','BR-08 timed for January 2027 tables; PL-08 before agricultural-tax season (Jan–Mar) with gmina→okręg data.','2027 parameter publication','Low–Medium'),
 ('Localisation expansions','After Batch 1 KPIs','BR-01 engine → ID/JP/ES/PL/TR; NUM → MX/ID/DE/ES/TR','Only expand where the lead market hits KPIs.','Batch 1 results','Low'),
],columns=['batch','timing','opportunities','scope_and_dependencies','depends_on','effort'])
validation=pd.DataFrame([
 ('Geo-located SERP snapshots','All Top 20 (esp. #13–#20)','Our SERP evidence is a US-index proxy; extended re-checks already overturned several "weak SERP" verdicts. Pull top-10 from google.com.br/.mx/.co.id/.co.jp/.pl/.de with a SERP API for each primary + 2 secondary keywords.','Could reorder #9–#20 and remove entries','High'),
 ('Authority/traffic of incumbents','calculadorarural.com.br, valorudi.com, calculadora-udi.com, kalenderindo.id, hitunganweton.id, meetly.com.pl, poocoo.pl, rvg-rechner.de, kalkulatory.gofin.pl, haradaeigo.com','Third-party DR/traffic (Ahrefs/Semrush/Similarweb) to separate weak EMDs from strong sites.','Changes A scores','High'),
 ('Precise monthly volumes & trends','Top-20 heads','Screenshots give 10× ranges; several heads show −90% YoY (verzugszinsen berechnen, hektar przeliczeniowy, calculadora modalidad 40, 1 mayam berapa gram, calculadora de liquidación). Pull 12-month series (any spending account or GKP export) to separate real decline from bucket artefacts.','Changes B scores; could drop DE-07/PL-08','High'),
 ('Turkey Keyword Planner export','TR (1,277 keywords, zero metrics)','No TR screenshot exists; TR cannot enter the Top 20 without it.','Could add 1–3 TR entries (dönüm, kira artış, 4\'lük)','Medium'),
 ('Non-Japan locations for English Japan queries','JP_EN','jp_en plan was measured with location = Japan, which excludes most English searchers. Re-run with Worldwide/US/PH/VN/IN.','Could promote JP_EN-02 pension refund (strong monetisation)','Medium'),
 ('Keywords not visible in screenshots','1,837 CSV keywords (18.4% of non-TR) incl. "uma 2026" family, "validador cnpj alfanumérico", many JP tails','Stitched captures skip ~7 rows between segments.','Could enlarge MX-03 and BR-04 clusters','Medium'),
 ('Legal/regulatory sign-off','MX-01, BR-03, BR-02, BR-08, PL-04, JP-07, DE-02/03/07','Formulas and parameters must be validated by a local professional before launch (e.g., MX prima de antigüedad cap basis; PL 2026 staż rules).','Launch blocker for those tools','High'),
 ('Data licences','PL-02 (SGJP/Morfeusz), PL-05 (word list), JP-01 (Japan Post), MX-03 (Banxico SIE), BR-03 (CNAE×LC116 correlation)','Confirm reuse terms.','Launch blocker','Medium'),
 ('Monetisation assumptions','ID-01, ID-09, JP-TXT, PL-05','RPM in Indonesia/for word games is low; if the business model is RPM-driven, swap these for MX-02/BR-02/PL-01.','Reorders #4–#9','Medium'),
],columns=['missing_evidence','applies_to','why_and_how','decision_impact','priority'])
# ---- write xlsx
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
readme=pd.DataFrame({'item':['Purpose','Generated','Evidence','Keyword Planner settings (all screenshots)','Currency','Extraction method','Verification','SERP method','Important caveats','Sheets'],
 'detail':['Top-20 international organic opportunities for a multi-country tools/calculator site.','2026-10-12',
 '9 CSVs (11,238 keywords; no metrics) + 22 Google Ads Keyword Planner screenshots (BR4, DE2, ES3, ID4, JP2, JP_EN1, MX3, PL3; none for TR).',
 'Location = market country (Japan for JP_EN), All languages, Google, Sep 2025 – Aug 2026, "Saved keywords" view. Values are BUCKETED ranges (account without spend) — never converted to point estimates.',
 'Top-of-page bids are in Indian Rupees (₹) — the Ads account currency, not local currency.',
 'Row detection on separator lines + per-column Tesseract OCR; rupee glyph blanked before OCR; keywords aligned to CSV (punctuation-normalised, order-aware).',
 '76 distinct metric rows (finalist heads and high-bid rows) checked visually against the screenshots: 76/76 identical to the OCR output.',
 'WebSearch tool, native-language queries, standard + extended modes; US-index proxy, NOT geo-located Google. Direct Google/Bing/DDG fetches blocked (403).',
 'Dashes and missing values are UNKNOWN, not zero. Ads competition ≠ organic difficulty. Cluster floor/ceiling sums add bucket bounds and may double-count close variants.',
 'Top20, Top20_Detail, Top20_Keyword_Clusters, All_Candidates_Scored, Scoring_Rubric, SERP_Evidence, Discovery_List, Roadmap, Validation_Checklist, Screenshot_Audit, KP_All_Keywords']})
kp=master.rename(columns={'vol':'avg_monthly_searches','m3':'three_month_change','yoy':'yoy_change','comp':'ads_competition','bid_lo':'top_bid_low_INR','bid_hi':'top_bid_high_INR'})
kp=kp[['market','country','language','source_file','keyword','in_screenshot','avg_monthly_searches','three_month_change','yoy_change','ads_competition','top_bid_low_INR','top_bid_high_INR','screenshot_files','validation_status','status','match_score','ocr_kw']]
kp['avg_monthly_searches']=kp.apply(lambda r: r.avg_monthly_searches if isinstance(r.avg_monthly_searches,str) and r.avg_monthly_searches else ('— (unknown)' if r.in_screenshot else ''),axis=1)
sheets=[('README',readme),('Top20',T20[['rank','id','short_name','concept','markets','language','primary_keyword','primary_avg_monthly_searches','primary_3m_change','primary_yoy','primary_ads_competition','primary_top_bid_INR','cluster_kw_with_volume','cluster_10k_plus','cluster_1k_10k','cluster_100_1k','cluster_floor_sum','A_organic_25','B_demand_20','C_toolfit_20','D_intent_10','E_cluster_10','F_monetisation_10','G_feasibility_5','total_100','confidence','organic_assessment','url_structure']]),
 ('Top20_Detail',T20),('Top20_Keyword_Clusters',mem),('All_Candidates_Scored',allc),('Scoring_Rubric',rubric),('SERP_Evidence',serp_df),('Discovery_List',disc),
 ('Roadmap',roadmap),('Validation_Checklist',validation),('Screenshot_Audit',audit),('KP_All_Keywords',kp)]
xl=os.path.join(OUT,'top20_seo_opportunities.xlsx')
with pd.ExcelWriter(xl,engine='openpyxl') as w:
    for name,df in sheets: df.to_excel(w,sheet_name=name,index=False)
    for name,df in sheets:
        ws=w.sheets[name]; ws.freeze_panes='A2'
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78'); cell.alignment=Alignment(wrap_text=True,vertical='top')
        for i,col in enumerate(df.columns,1):
            mx=max([len(str(col))]+[len(str(v)) for v in df[col].head(200)])
            ws.column_dimensions[get_column_letter(i)].width=min(max(10,mx*0.9),60)
        if len(df)<400:
            for row in ws.iter_rows(min_row=2):
                for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
        ws.auto_filter.ref=ws.dimensions
# csv copies
T20.to_csv(os.path.join(OUT,'top20.csv'),index=False); allc.to_csv(os.path.join(OUT,'candidates_scored.csv'),index=False)
mem.to_csv(os.path.join(OUT,'top20_keyword_clusters.csv'),index=False); kp.to_csv(os.path.join(OUT,'keyword_metrics_extracted.csv'),index=False)
serp_df.to_csv(os.path.join(OUT,'serp_evidence.csv'),index=False); disc.to_csv(os.path.join(OUT,'discovery_list.csv'),index=False)
audit.to_csv(os.path.join(OUT,'screenshot_audit.csv'),index=False)
json.dump(dict(cands=[{k:v for k,v in c.items() if k!='spec'} for c in cands]),open(os.path.join(S,'cands_computed.json'),'w'),ensure_ascii=False,default=str,indent=1)
print('written',xl)
print(T20[['rank','id','primary_keyword','primary_avg_monthly_searches','cluster_kw_with_volume','cluster_10k_plus','cluster_1k_10k','cluster_100_1k','cluster_floor_sum','total_100','max_top_bid_high_INR']].to_string())
print(len(mem), mem.validation_status.str.startswith('NOT').sum())
